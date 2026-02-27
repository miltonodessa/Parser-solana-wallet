"""
DataExporter: saves decoded transaction records and analysis results to disk.

Supported formats:
  • JSON  – full structured data (nested dicts / lists preserved)
  • CSV   – flat per-transaction table
  • XLSX  – multi-sheet Excel workbook (transactions + analysis sheets)
"""
from __future__ import annotations

import csv
import json
import logging
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

logger = logging.getLogger(__name__)


class DataExporter:
    """Writes parsed Solana data to various file formats."""

    def __init__(self, output_dir: str = "output") -> None:
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    # ── JSON ───────────────────────────────────────────────────────────────────

    def to_json(
        self,
        data: Any,
        filename: str,
        indent: int = 2,
    ) -> Path:
        """Write *data* as pretty-printed JSON."""
        path = self.output_dir / filename
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=indent, default=str)
        logger.info("Saved JSON → %s", path)
        return path

    def transactions_to_json(
        self,
        records: list[dict],
        wallet: str,
        filename: Optional[str] = None,
    ) -> Path:
        filename = filename or f"{_safe(wallet)}_transactions.json"
        payload = {
            "wallet": wallet,
            "exported_at": datetime.now(timezone.utc).isoformat(),
            "total_transactions": len(records),
            "transactions": records,
        }
        return self.to_json(payload, filename)

    def analysis_to_json(
        self,
        analysis: dict,
        wallet: str,
        filename: Optional[str] = None,
    ) -> Path:
        filename = filename or f"{_safe(wallet)}_analysis.json"
        return self.to_json(analysis, filename)

    # ── CSV ────────────────────────────────────────────────────────────────────

    def transactions_to_csv(
        self,
        records: list[dict],
        wallet: str,
        filename: Optional[str] = None,
    ) -> Path:
        """
        Write a flat CSV with one row per transaction.

        Nested structures (instructions, token_changes, etc.) are serialised
        as JSON strings so nothing is lost.
        """
        filename = filename or f"{_safe(wallet)}_transactions.csv"
        path = self.output_dir / filename

        flat_records = [_flatten_record(r) for r in records]
        if not flat_records:
            logger.warning("No records to write to CSV")
            path.write_text("", encoding="utf-8")
            return path

        fieldnames = list(flat_records[0].keys())
        # Ensure core fields are first
        priority = [
            "signature", "block_time_utc", "slot", "success", "error",
            "fee_sol", "fee_lamports", "fee_payer", "version",
        ]
        ordered = [f for f in priority if f in fieldnames]
        ordered += [f for f in fieldnames if f not in ordered]

        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=ordered, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(flat_records)

        logger.info("Saved CSV (%d rows) → %s", len(flat_records), path)
        return path

    def sol_transfers_to_csv(
        self,
        transfers: list[dict],
        wallet: str,
        filename: Optional[str] = None,
    ) -> Path:
        filename = filename or f"{_safe(wallet)}_sol_transfers.csv"
        return self._list_to_csv(transfers, filename)

    def token_transfers_to_csv(
        self,
        transfers: list[dict],
        wallet: str,
        filename: Optional[str] = None,
    ) -> Path:
        filename = filename or f"{_safe(wallet)}_token_transfers.csv"
        return self._list_to_csv(transfers, filename)

    def swaps_to_csv(
        self,
        swaps: list[dict],
        wallet: str,
        filename: Optional[str] = None,
    ) -> Path:
        filename = filename or f"{_safe(wallet)}_swaps.csv"
        flat = [_flatten_record(s) for s in swaps]
        return self._list_to_csv(flat, filename)

    # ── XLSX ───────────────────────────────────────────────────────────────────

    def to_xlsx(
        self,
        records: list[dict],
        analysis: dict,
        wallet: str,
        filename: Optional[str] = None,
    ) -> Optional[Path]:
        """
        Write a multi-sheet Excel workbook.

        Requires the `openpyxl` package (optional dependency).
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            logger.warning("openpyxl not installed – skipping XLSX export. Run: pip install openpyxl")
            return None

        filename = filename or f"{_safe(wallet)}_full_report.xlsx"
        path = self.output_dir / filename
        wb = openpyxl.Workbook()

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1B4F8A")
        header_align = Alignment(horizontal="center")

        def _write_sheet(ws, rows: list[dict]) -> None:
            if not rows:
                return
            flat = [_flatten_record(r) for r in rows]
            cols = list(flat[0].keys())
            ws.append(cols)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
            for row in flat:
                ws.append([row.get(c) for c in cols])
            # Auto-width
            for i, col in enumerate(cols, 1):
                max_len = max(len(str(col)), *(len(str(r.get(col, "") or "")) for r in flat))
                ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 60)

        # Sheet 1: Transactions
        ws_tx = wb.active
        ws_tx.title = "Transactions"
        _write_sheet(ws_tx, records)

        # Sheet 2: SOL Transfers
        sol_t = analysis.get("sol_transfers", [])
        if sol_t:
            _write_sheet(wb.create_sheet("SOL Transfers"), sol_t)

        # Sheet 3: Token Transfers
        tok_t = analysis.get("token_transfers", [])
        if tok_t:
            _write_sheet(wb.create_sheet("Token Transfers"), tok_t)

        # Sheet 4: Swaps
        swaps = analysis.get("swaps", [])
        if swaps:
            flat_swaps = [_flatten_record(s) for s in swaps]
            _write_sheet(wb.create_sheet("Swaps"), flat_swaps)

        # Sheet 5: NFT Events
        nft = analysis.get("nft_events", [])
        if nft:
            flat_nft = [_flatten_record(n) for n in nft]
            _write_sheet(wb.create_sheet("NFT Events"), flat_nft)

        # Sheet 6: Summary
        summary = analysis.get("summary", {})
        ws_sum = wb.create_sheet("Summary")
        ws_sum.append(["Metric", "Value"])
        for cell in ws_sum[1]:
            cell.font = header_font
            cell.fill = header_fill
        for k, v in summary.items():
            ws_sum.append([k, str(v)])
        ws_sum.column_dimensions["A"].width = 35
        ws_sum.column_dimensions["B"].width = 35

        # Sheet 7: Programs Used
        prog = analysis.get("programs_used", {})
        ws_prog = wb.create_sheet("Programs Used")
        ws_prog.append(["Program", "Invocations"])
        for cell in ws_prog[1]:
            cell.font = header_font
            cell.fill = header_fill
        for prog_name, count in prog.items():
            ws_prog.append([prog_name, count])

        wb.save(path)
        logger.info("Saved XLSX → %s", path)
        return path

    # ── Internal helpers ───────────────────────────────────────────────────────

    def _list_to_csv(self, rows: list[dict], filename: str) -> Path:
        path = self.output_dir / filename
        if not rows:
            path.write_text("", encoding="utf-8")
            return path
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        logger.info("Saved CSV (%d rows) → %s", len(rows), path)
        return path


# ─── Utilities ─────────────────────────────────────────────────────────────────

def _safe(name: str) -> str:
    """Make a string safe for use as a filename."""
    return name[:20].replace("/", "_").replace("\\", "_")


def _flatten_record(record: dict) -> dict:
    """
    Flatten a nested transaction record for CSV/XLSX output.

    Nested lists/dicts are JSON-serialised as strings.
    """
    flat: dict[str, Any] = {}
    for k, v in record.items():
        if isinstance(v, (dict, list)):
            flat[k] = json.dumps(v, ensure_ascii=False, default=str)
        else:
            flat[k] = v
    return flat
